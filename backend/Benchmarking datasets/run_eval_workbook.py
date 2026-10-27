"""Run every validator over guardrails_validator_eval_set.xlsx and record one run.

For each validator sheet it builds the validator from the Configs sheet (via the
same config classes the API uses), calls validate() on input_text, and writes:

- runs/<date>/eval_run_<date>.xlsx  workbook copy with actual_result / actual_output filled
- runs/<date>/breakdown.json         metrics per validator, per category and per language
- Run_Log rows appended to guardrails_validator_eval_set.xlsx

Usage (from backend/):
    uv run --with openpyxl python "Benchmarking datasets/run_eval_workbook.py" [validator ...]
"""

import json
import re
import subprocess
import sys
import time
import traceback
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path

from guardrails.validators import FailResult
from openpyxl import load_workbook

from app.core.validators.config.answer_relevance_custom_llm_safety_validator_config import (
    AnswerRelevanceCustomLLMSafetyValidatorConfig,
)
from app.core.validators.config.ban_list_safety_validator_config import (
    BanListSafetyValidatorConfig,
)
from app.core.validators.config.content_safety_safety_validator_config import (
    ContentSafetySafetyValidatorConfig,
)
from app.core.validators.config.gender_assumption_bias_safety_validator_config import (
    GenderAssumptionBiasSafetyValidatorConfig,
)
from app.core.validators.config.lexical_slur_safety_validator_config import (
    LexicalSlurSafetyValidatorConfig,
)
from app.core.validators.config.llamaguard_7b_safety_validator_config import (
    LlamaGuard7BSafetyValidatorConfig,
)
from app.core.validators.config.llm_critic_safety_validator_config import (
    LLMCriticSafetyValidatorConfig,
)
from app.core.validators.config.nsfw_text_safety_validator_config import (
    NSFWTextSafetyValidatorConfig,
)
from app.core.validators.config.pii_remover_safety_validator_config import (
    PIIRemoverSafetyValidatorConfig,
)
from app.core.validators.config.profanity_free_safety_validator_config import (
    ProfanityFreeSafetyValidatorConfig,
)
from app.core.validators.config.topic_relevance_llm_safety_validator_config import (
    TopicRelevanceLLMSafetyValidatorConfig,
)
from app.core.validators.config.topic_relevance_safety_validator_config import (
    TopicRelevanceSafetyValidatorConfig,
)

ROOT = Path(__file__).resolve().parent
WORKBOOK = ROOT / "guardrails_validator_eval_set.xlsx"

CONFIG_CLASSES = {
    "content_safety": ContentSafetySafetyValidatorConfig,
    "uli_slur_match": LexicalSlurSafetyValidatorConfig,
    "profanity_free": ProfanityFreeSafetyValidatorConfig,
    "nsfw_text": NSFWTextSafetyValidatorConfig,
    "llamaguard_7b": LlamaGuard7BSafetyValidatorConfig,
    "pii_remover": PIIRemoverSafetyValidatorConfig,
    "gender_assumption_bias": GenderAssumptionBiasSafetyValidatorConfig,
    "topic_relevance_llm": TopicRelevanceLLMSafetyValidatorConfig,
    "topic_relevance": TopicRelevanceSafetyValidatorConfig,
    "ban_list": BanListSafetyValidatorConfig,
    "answer_relevance_custom_llm": AnswerRelevanceCustomLLMSafetyValidatorConfig,
    "llm_critic": LLMCriticSafetyValidatorConfig,
}
# Sheet name -> validator type, where they differ.
SHEET_VALIDATOR = {"answer_relevance": "answer_relevance_custom_llm"}
LLM_VALIDATORS = {"topic_relevance_llm", "topic_relevance", "answer_relevance_custom_llm", "llm_critic"}
LLM_WORKERS = 8


def git_sha() -> str:
    try:
        return subprocess.check_output(["git", "rev-parse", "--short", "HEAD"], cwd=ROOT, text=True).strip()
    except Exception:
        return ""


def load_configs(wb) -> dict:
    ws = wb["Configs"]
    configs = {}
    for ref, validator, params, *_ in ws.iter_rows(min_row=4, values_only=True):
        if ref:
            configs[ref] = (validator, json.loads(params))
    return configs


def build_validator(validator_type: str, params: dict, row: dict):
    params = {k: v for k, v in params.items() if v is not None}
    if validator_type == "answer_relevance_custom_llm":
        params = {**params, "input": row["query"], "output": row["answer"]}
    return CONFIG_CLASSES[validator_type](type=validator_type, **params).build()


def run_case(validator, text: str):
    start = time.perf_counter()
    try:
        result = validator.validate(text, {})
        elapsed = (time.perf_counter() - start) * 1000
        if isinstance(result, FailResult):
            return "FAIL", result.fix_value or "", result.error_message or "", elapsed
        return "PASS", text, "", elapsed
    except Exception as e:
        elapsed = (time.perf_counter() - start) * 1000
        return "ERROR", "", f"{type(e).__name__}: {e}", elapsed


def binary_metrics(pairs):
    """pairs: list of (expected, actual) with actual in PASS/FAIL/ERROR."""
    tp = sum(e == "FAIL" and a == "FAIL" for e, a in pairs)
    fp = sum(e == "PASS" and a == "FAIL" for e, a in pairs)
    fn = sum(e == "FAIL" and a == "PASS" for e, a in pairs)
    tn = sum(e == "PASS" and a == "PASS" for e, a in pairs)
    errors = sum(a == "ERROR" for _, a in pairs)
    n = tp + fp + fn + tn
    precision = tp / (tp + fp) if tp + fp else None
    recall = tp / (tp + fn) if tp + fn else None
    f1 = 2 * precision * recall / (precision + recall) if precision and recall else (0.0 if precision is not None and recall is not None else None)
    return {
        "rows_run": n, "errors": errors, "tp": tp, "fp": fp, "fn": fn, "tn": tn,
        "accuracy": (tp + tn) / n if n else None,
        "precision": precision, "recall": recall, "f1": f1,
    }


def percentile(values, q):
    if not values:
        return None
    s = sorted(values)
    return s[min(len(s) - 1, int(round(q * (len(s) - 1))))]


def norm_ws(s: str) -> str:
    s = " ".join(str(s).split())
    return re.sub(r"\s+([?.,!।])", r"\1", s)


def run_sheet(ws, configs):
    header = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(header)}
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[col["case_id"]]]
    validator_type = SHEET_VALIDATOR.get(ws.title, ws.title)

    # One validator per config_ref; answer relevance needs one per row.
    cache = {}

    def get_validator(row):
        if validator_type == "answer_relevance_custom_llm":
            _, params = configs[row["config_ref"]]
            return build_validator(validator_type, params, row)
        ref = row["config_ref"]
        if ref not in cache:
            _, params = configs[ref]
            cache[ref] = build_validator(validator_type, params, row)
        return cache[ref]

    def task(row):
        try:
            validator = get_validator(row)
        except Exception as e:
            return "ERROR", "", f"build failed: {type(e).__name__}: {e}", 0.0
        return run_case(validator, str(row["input_text"]))

    if validator_type in LLM_VALIDATORS:
        for row in rows[:1]:  # build once up front so per-thread builds hit the cache
            if validator_type != "answer_relevance_custom_llm":
                try:
                    get_validator(row)
                except Exception:
                    pass
        with ThreadPoolExecutor(LLM_WORKERS) as pool:
            results = list(pool.map(task, rows))
    else:
        results = [task(r) for r in rows]
    return validator_type, rows, results, col


def summarize(validator_type, rows, results):
    pairs = [(r["expected_result"], res[0]) for r, res in zip(rows, results)]
    overall = binary_metrics(pairs)
    latencies = [res[3] for res in results if res[0] != "ERROR"]
    overall["latency_p50_ms"] = percentile(latencies, 0.5)
    overall["latency_p95_ms"] = percentile(latencies, 0.95)

    fix_rows = [
        (r, res) for r, res in zip(rows, results)
        if r["expected_result"] == "FAIL" and r.get("expected_output") and res[0] == "FAIL"
    ]
    if fix_rows and validator_type != "pii_remover":  # PII targets use dataset placeholders, not Presidio's
        overall["fix_exact_match"] = sum(norm_ws(res[1]) == norm_ws(r["expected_output"]) for r, res in fix_rows) / len(fix_rows)
        overall["fix_rows_compared"] = len(fix_rows)

    breakdown = {}
    for key in ("category", "language", "config_ref", "source"):
        groups = defaultdict(list)
        for r, (e, a) in zip(rows, pairs):
            groups[str(r.get(key) or "")].append((e, a))
        breakdown[key] = {g: binary_metrics(p) for g, p in sorted(groups.items())}

    errors = defaultdict(int)
    for res in results:
        if res[0] == "ERROR":
            errors[res[2][:160]] += 1
    misses = [
        {"case_id": r["case_id"], "expected": r["expected_result"], "actual": res[0],
         "category": r.get("category"), "language": r.get("language"),
         "input_text": str(r["input_text"])[:200], "detail": res[2][:200]}
        for r, res in zip(rows, results) if res[0] in ("PASS", "FAIL") and res[0] != r["expected_result"]
    ]
    return {"overall": overall, "breakdown": breakdown, "error_messages": dict(errors), "misclassified": misses}


def append_run_log(ws, entries):
    r = 5
    # Drop the template's example row; then append after the last filled row.
    if str(ws.cell(5, 14).value or "").startswith("EXAMPLE"):
        for c in range(1, 15):
            ws.cell(5, c).value = None
    while ws.cell(r, 1).value:
        r += 1
    for e in entries:
        for c, v in enumerate(e, 1):
            ws.cell(r, c).value = v
        r += 1


def main(only: list[str]):
    today = date.today().isoformat()
    out_dir = ROOT / "runs" / today
    out_dir.mkdir(parents=True, exist_ok=True)
    sha = git_sha()

    run_copy = out_dir / f"eval_run_{today}.xlsx"
    # A partial run (some validators only) builds on today's copy if one exists.
    wb = load_workbook(run_copy if run_copy.exists() else WORKBOOK)
    configs = load_configs(wb)
    report, log_entries = {}, []
    for sheet in [s for s in wb.sheetnames if SHEET_VALIDATOR.get(s, s) in CONFIG_CLASSES]:
        validator_type = SHEET_VALIDATOR.get(sheet, sheet)
        if only and validator_type not in only:
            continue
        ws = wb[sheet]
        print(f"running {validator_type} ...", flush=True)
        t0 = time.time()
        try:
            validator_type, rows, results, col = run_sheet(ws, configs)
        except Exception:
            traceback.print_exc()
            continue
        for i, res in enumerate(results, start=2):
            ws.cell(i, col["actual_result"] + 1).value = res[0]
            ws.cell(i, col["actual_output"] + 1).value = res[1] if res[0] != "ERROR" else res[2]
        summary = summarize(validator_type, rows, results)
        report[validator_type] = summary
        m = summary["overall"]
        print(f"  {time.time() - t0:.0f}s  " + json.dumps({k: (round(v, 3) if isinstance(v, float) else v) for k, v in m.items()}), flush=True)

        params = configs.get(rows[0]["config_ref"], (None, {}))[1] if rows else {}
        model = params.get("llm_callable") or params.get("model_name") or ""
        notes = []
        if m["errors"]:
            notes.append(f"{m['errors']} errors: {next(iter(summary['error_messages']))[:120]}")
        if "fix_exact_match" in m:
            notes.append(f"fix exact match {m['fix_exact_match']:.2f} on {m['fix_rows_compared']} rows")
        rnd = lambda v: round(v, 3) if isinstance(v, float) else v
        log_entries.append([
            today, sha, validator_type, ", ".join(sorted({r["config_ref"] for r in rows})), model,
            m["rows_run"], rnd(m["accuracy"]), rnd(m["precision"]), rnd(m["recall"]), rnd(m["f1"]),
            rnd(m["latency_p50_ms"]), rnd(m["latency_p95_ms"]), "claude-code", "; ".join(notes),
        ])

    append_run_log(wb["Run_Log"], log_entries)
    wb.save(run_copy)
    breakdown_path = out_dir / "breakdown.json"
    if breakdown_path.exists():
        report = {**json.loads(breakdown_path.read_text()), **report}
    breakdown_path.write_text(json.dumps(report, indent=2, ensure_ascii=False, default=str))

    main_wb = load_workbook(WORKBOOK)
    append_run_log(main_wb["Run_Log"], log_entries)
    main_wb.save(WORKBOOK)
    print(f"saved {run_copy}\nappended {len(log_entries)} Run_Log rows to {WORKBOOK.name}")


if __name__ == "__main__":
    main(sys.argv[1:])
