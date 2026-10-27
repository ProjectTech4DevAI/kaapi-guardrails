"""Build the combined guardrails validator evaluation workbook.

Compiles the benchmark CSVs (Content Safety, Privacy protection, Gender neutrality,
Scope Control) and the synthetic CSVs in `Synthetic/` into one XLSX with a common
schema, a config registry, a live metrics index and a run log.

Usage (from this directory):
    uv run --with openpyxl --with pandas python build_eval_workbook.py
"""

import json
import re
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
SYNTH = ROOT / "Synthetic"
OUT = ROOT / "guardrails_validator_eval_set.xlsx"

CORE_COLS = [
    "case_id", "validator", "input_text", "expected_result", "label",
    "expected_output", "category", "difficulty", "language", "domain",
    "config_ref", "source", "notes",
]
RUN_COLS = ["actual_result", "actual_output", "match"]

DEVANAGARI = re.compile(r"[ऀ-ॿ]")
HINGLISH_MARKERS = set(
    "kya hai hain kab kaise kaisa kaisi ke ki ka mein me karna karne chahiye kyon kyun "
    "doon du nahi nahin ko se aur kitna kitni liye bacche baccha bachcha kar hota hoti "
    "mera meri mere kaun kisne kahan wala wali hum tum aap apna apni kuch bhi sakte "
    "sakta sakti batao bataye bata raha rahi tha thi hoga milega milta yojana garbh "
    "khana dena lena jaye jaana padhai likha likhe walo haar".split()
)


def detect_lang(text: str) -> str:
    """Heuristic language tag: hi (Devanagari) / hi-Latn (romanized Hindi) / en."""
    if DEVANAGARI.search(text):
        return "hi"
    words = re.findall(r"[a-z]+", text.lower())
    hits = sum(w in HINGLISH_MARKERS for w in words)
    return "hi-Latn" if words and hits / len(words) >= 0.2 or hits >= 2 else "en"


def finalize(df: pd.DataFrame, validator: str, prefix: str) -> pd.DataFrame:
    df = df.copy()
    df["validator"] = validator
    df["expected_result"] = df["expected_result"].str.upper().str.strip()
    df["label"] = (df["expected_result"] == "FAIL").astype(int)
    for c in CORE_COLS:
        if c not in df.columns:
            df[c] = ""
    df = df.fillna("")
    df["case_id"] = [f"{prefix}-{i:04d}" for i in range(1, len(df) + 1)]
    extras = [c for c in df.columns if c not in CORE_COLS]
    return df[CORE_COLS + extras]


# --------------------------------------------------------------------------- benchmark loaders

def load_content_safety():
    df = pd.read_csv(ROOT / "Content Safety/toxicity_test_combined.csv")
    lang = df["language"].str.lower().map({"english": "en", "en": "en", "hindi": "hi", "hi": "hi"})
    is_deva = df["text"].astype(str).str.contains(DEVANAGARI)
    lang = lang.where(~((lang == "hi") & ~is_deva), "hi-Latn")
    out = pd.DataFrame({
        "input_text": df["text"],
        "expected_result": df["label"].map({1: "FAIL", 0: "PASS"}),
        "category": df["label"].map({1: "toxic", 0: "non_toxic"}),
        "language": lang,
        "domain": "general",
        "config_ref": "CS-DEFAULT",
        "source": "benchmark:" + df["dataset"],
        "source_dataset": df["dataset"],
    })
    return finalize(out, "content_safety", "CS")


def load_pii():
    df = pd.read_csv(ROOT / "Privacy protection/pii_detection_testing_dataset.csv")
    entities = df["target_text"].fillna("").apply(
        lambda t: ";".join(sorted(set(re.findall(r"\[([A-Z_]+)\]", t))))
    )
    out = pd.DataFrame({
        "input_text": df["source_text"],
        "expected_result": df["label"].map({"pii": "FAIL", "no_pii": "PASS"}),
        "expected_output": df["target_text"].fillna(df["source_text"]),
        "category": df["label"],
        "language": df["language"],
        "domain": "general",
        "config_ref": "PII-DEFAULT",
        "source": "benchmark:pii_detection_testing_dataset",
        "pii_entities": entities.replace("", "none"),
        "entity_count": df["target_text"].fillna("").str.count(r"\[[A-Z_]+\]"),
    })
    return finalize(out, "pii_remover", "PII")


def load_gender():
    df = pd.read_csv(ROOT / "Gender neutrality/gender_bias_assumption_dataset.csv")
    common = {
        "language": df["language tag"].map({"English": "en", "Hindi": "hi", "Hinglish": "hi-Latn"}),
        "domain": df["domain"],
        "config_ref": "GAB-ALL",
        "source": "benchmark:gender_bias_assumption_dataset",
        "transformation_type": df["transformation type"],
    }
    # Same pairing as app/evaluation/gender_assumption_bias/run.py: the biased
    # sentence should FAIL and its neutral rewrite should PASS unchanged.
    biased = pd.DataFrame({
        "input_text": df["biased input"], "expected_result": "FAIL",
        "expected_output": df["neutral output"], "category": df["bias type"], **common,
    })
    neutral = pd.DataFrame({
        "input_text": df["neutral output"], "expected_result": "PASS",
        "expected_output": df["neutral output"], "category": "already_neutral", **common,
    })
    out = pd.concat([biased, neutral], ignore_index=True)
    out["notes"] = out.duplicated("input_text").map(
        {True: "duplicate input_text (dataset repeats 15 templates x20)", False: ""}
    )
    return finalize(out, "gender_assumption_bias", "GAB")


def load_scope():
    parts = []
    for domain, ref in [("education", "TR-EDU"), ("healthcare", "TR-HEALTH")]:
        df = pd.read_csv(ROOT / f"Scope Control/{domain}-topic-relevance-dataset.csv")
        parts.append(pd.DataFrame({
            "input_text": df["input"],
            "expected_result": df["scope"].map({"OUT_OF_SCOPE": "FAIL", "IN_SCOPE": "PASS"}),
            "category": df["category"].str.strip(),
            "language": df["input"].astype(str).apply(detect_lang),
            "domain": domain,
            "config_ref": ref,
            "source": f"benchmark:{domain}-topic-relevance-dataset",
            "scope": df["scope"],
        }))
    return finalize(pd.concat(parts, ignore_index=True), "topic_relevance_llm", "TRL")


def load_synthetic(validator: str, prefix: str):
    df = pd.read_csv(SYNTH / f"{validator}.csv", dtype=str, keep_default_na=False)
    df["source"] = "synthetic"
    return finalize(df, validator, prefix)


# --------------------------------------------------------------------------- metadata

# (validator, sheet, prefix, group, origin, loader, positive-class meaning, primary metric, stage, runner)
VALIDATORS = [
    ("content_safety", "content_safety", "CS", "Content safety", "Benchmark", load_content_safety,
     "Text is toxic/abusive", "Precision / Recall / F1", "input + output", "app/evaluation/toxicity/run.py"),
    ("uli_slur_match", "uli_slur_match", "SLUR", "Content safety", "Synthetic", None,
     "Text contains a listed slur (abusive intent)", "Recall; exact match of redacted text", "input + output", "app/evaluation/toxicity/run.py"),
    ("profanity_free", "profanity_free", "PROF", "Content safety", "Synthetic", None,
     "Text contains profanity", "Precision / Recall / F1", "input + output", "app/evaluation/toxicity/run.py"),
    ("nsfw_text", "nsfw_text", "NSFW", "Content safety", "Synthetic", None,
     "Text is sexually explicit / NSFW", "Precision (false positives on health Qs)", "input + output", "app/evaluation/toxicity/run.py"),
    ("llamaguard_7b", "llamaguard_7b", "LG", "Content safety", "Synthetic", None,
     "Text violates a LlamaGuard policy", "Recall per policy; FPR on look-alikes", "input + output", "app/evaluation/toxicity/run.py"),
    ("pii_remover", "pii_remover", "PII", "Privacy", "Benchmark", load_pii,
     "Text contains PII", "Precision / Recall / F1; per-entity", "input + output", "app/evaluation/pii/run.py"),
    ("gender_assumption_bias", "gender_assumption_bias", "GAB", "Bias", "Benchmark", load_gender,
     "Text contains gender-assumptive language", "Recall; exact match of neutral rewrite", "output", "app/evaluation/gender_assumption_bias/run.py"),
    ("topic_relevance_llm", "topic_relevance_llm", "TRL", "Scope control", "Benchmark", load_scope,
     "Message is out of scope", "Accuracy / F1; in-scope wrongly blocked", "input", "app/evaluation/topic_relevance/run.py"),
    ("topic_relevance", "topic_relevance", "TR", "Scope control", "Synthetic", None,
     "Message is out of scope", "Accuracy / F1", "input", "app/evaluation/topic_relevance/run.py"),
    ("ban_list", "ban_list", "BAN", "Scope control", "Synthetic", None,
     "Text contains a banned term (policy intent)", "Recall; exact match of fixed text", "input + output", "app/evaluation/ban_list/run.py"),
    ("answer_relevance_custom_llm", "answer_relevance", "AR", "Response quality", "Synthetic", None,
     "Answer does not satisfy the query", "Accuracy; correct answers wrongly rejected", "output", "— (no runner yet)"),
    ("llm_critic", "llm_critic", "CRIT", "Response quality", "Synthetic", None,
     "Bot response falls below a metric threshold", "Accuracy; per-metric agreement", "output", "— (no runner yet)"),
]


def read_text(p: Path) -> str:
    return p.read_text(encoding="utf-8").strip() if p.exists() else ""


def build_configs():
    critic_cfg = SYNTH / "llm_critic_config.json"
    critic = json.loads(critic_cfg.read_text()) if critic_cfg.exists() else {}
    rows = [
        ("CS-DEFAULT", "content_safety", {"on_fail": "fix"}, "Slur redaction then alt-profanity-check scoring (default threshold).", ""),
        ("SLUR-DEFAULT", "uli_slur_match", {"languages": ["en", "hi"], "severity": "all", "on_fail": "fix"}, "Curated Hindi/English slur list, all severities.", "app/core/validators/utils/files/curated_slurlist_hi_en.csv"),
        ("PROF-DEFAULT", "profanity_free", {"on_fail": "exception"}, "alt-profanity-check linear SVM; no fix value.", ""),
        ("NSFW-DEFAULT", "nsfw_text", {"threshold": 0.8, "validation_method": "sentence", "model_name": "michellejieli/NSFW_text_classifier", "on_fail": "exception"}, "Sentence-level NSFW classifier.", ""),
        ("LLAMAGUARD-ALL", "llamaguard_7b", {"policies": None, "on_fail": "exception"}, "All six LlamaGuard policies enabled (remote inference).", ""),
        ("PII-DEFAULT", "pii_remover", {"on_fail": "fix"}, "Default pattern entity types; expected_output uses ai4privacy-style [ENTITY] placeholders.", ""),
        ("GAB-ALL", "gender_assumption_bias", {"categories": ["all"], "on_fail": "fix"}, "All bias categories.", "app/core/validators/utils/files/gender_assumption_bias_words.csv"),
        ("TR-EDU", "topic_relevance_llm", {"prompt_schema_version": 2, "threshold": 2, "llm_callable": "gpt-5.6-sol", "on_fail": "rephrase", "configuration": read_text(ROOT / "Scope Control/education_topic_config.txt")}, "Education bot — forbidden-topics config (v2).", "Scope Control/education_topic_config.txt"),
        ("TR-HEALTH", "topic_relevance_llm", {"prompt_schema_version": 2, "threshold": 2, "llm_callable": "gpt-5.6-sol", "on_fail": "rephrase", "configuration": read_text(ROOT / "Scope Control/healthcare_topic_config.txt")}, "Maternal health bot — forbidden-topics config (v2).", "Scope Control/healthcare_topic_config.txt"),
        ("TR-AGRI", "topic_relevance", {"prompt_schema_version": 1, "llm_callable": "gpt-5.6-sol", "on_fail": "rephrase", "configuration": read_text(SYNTH / "topic_relevance_agri_config.txt")}, "Agriculture advisory bot — allowed-topics config (v1).", "Synthetic/topic_relevance_agri_config.txt"),
        ("BAN-MATERNAL", "ban_list", {"banned_words": ["sonography", "gender check"], "on_fail": "fix"}, "Maternal health: sex-determination terms (PCPNDT Act).", "app/evaluation/ban_list/run.py"),
        ("BAN-EDU", "ban_list", {"banned_words": ["leaked paper", "proxy attendance", "fake certificate"], "on_fail": "fix"}, "Education: malpractice terms.", ""),
        ("AR-DEFAULT", "answer_relevance_custom_llm", {"llm_callable": "gpt-5.6-sol", "prompt_template": None, "on_fail": "exception"}, "Built-in strict YES/NO prompt. input_text is JSON {query, answer}.", ""),
        ("CRITIC-NGO", "llm_critic", critic or {"see": "Synthetic/llm_critic_config.json"}, "Simplicity / empathy / safety rubric for bot responses.", "Synthetic/llm_critic_config.json"),
    ]
    return rows


# --------------------------------------------------------------------------- styling helpers

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
RUN_FILL = PatternFill("solid", fgColor="FFF2CC")
RUN_HDR_FILL = PatternFill("solid", fgColor="BF8F00")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
BOLD = Font(name=FONT, bold=True)
BASE = Font(name=FONT, size=10)
HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

WIDTHS = {
    "case_id": 11, "validator": 18, "input_text": 60, "expected_result": 10, "label": 6,
    "expected_output": 50, "category": 24, "difficulty": 10, "language": 9, "domain": 13,
    "config_ref": 15, "source": 26, "notes": 40, "actual_result": 12, "actual_output": 40,
    "match": 8, "query": 40, "answer": 50,
}
WRAP_COLS = {"input_text", "expected_output", "notes", "query", "answer", "actual_output"}


def style_header(ws, row, ncols, fill=HDR_FILL):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def write_validator_sheet(wb, sheet, df):
    ws = wb.create_sheet(sheet)
    cols = list(df.columns) + RUN_COLS
    ws.append(cols)
    style_header(ws, 1, len(cols))
    for c in range(len(cols) - len(RUN_COLS) + 1, len(cols) + 1):
        ws.cell(row=1, column=c).fill = RUN_HDR_FILL

    idx = {c: i + 1 for i, c in enumerate(cols)}
    exp_col = get_column_letter(idx["expected_result"])
    act_col = get_column_letter(idx["actual_result"])
    n = len(df)
    for r, rec in enumerate(df.itertuples(index=False), start=2):
        ws.append(list(rec) + ["", "", f'=IF({act_col}{r}="","",IF({act_col}{r}={exp_col}{r},"Y","N"))'])

    for name, i in idx.items():
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = WIDTHS.get(name, 16)
        align = WRAP if name in WRAP_COLS else TOP
        for (cell,) in ws.iter_rows(min_row=2, max_row=n + 1, min_col=i, max_col=i):
            cell.font = BASE
            cell.alignment = align
            if name in RUN_COLS[:2]:
                cell.fill = RUN_FILL

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n + 1}"
    ws.row_dimensions[1].height = 30

    dv = DataValidation(type="list", formula1='"PASS,FAIL,ERROR"', allow_blank=True)
    dv.add(f"{act_col}2:{act_col}{n + 1}")
    ws.add_data_validation(dv)

    m = get_column_letter(idx["match"])
    rng = f"{m}2:{m}{n + 1}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Y"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"N"'], fill=PatternFill("solid", fgColor="FFC7CE")))
    e = f"{exp_col}2:{exp_col}{n + 1}"
    ws.conditional_formatting.add(e, FormulaRule(formula=[f'{exp_col}2="FAIL"'], font=Font(name=FONT, color="C00000", bold=True)))
    ws.conditional_formatting.add(e, FormulaRule(formula=[f'{exp_col}2="PASS"'], font=Font(name=FONT, color="2E7D32", bold=True)))
    return idx, n


def write_index(wb, meta):
    ws = wb.create_sheet("Index", 1)
    ws["A1"] = "Validator index & live metrics"
    ws["A1"].font = TITLE
    ws["A2"] = ("Counts and metrics are formulas over each validator sheet. Positive class = expected_result FAIL "
                "(validator should trigger). Metrics fill in once actual_result is entered on the validator sheets.")
    ws["A2"].font = Font(name=FONT, italic=True, size=9)
    headers = [
        "#", "validator", "sheet", "group", "data origin", "positive class (FAIL) means", "recommended stage",
        "primary metric", "eval runner (backend/)", "rows", "FAIL", "PASS", "% FAIL", "en", "hi", "hi-Latn",
        "configs", "rows run", "TP", "FP", "FN", "TN", "errors", "accuracy", "precision", "recall", "F1",
    ]
    hr = 4
    for c, h in enumerate(headers, 1):
        ws.cell(row=hr, column=c, value=h)
    style_header(ws, hr, len(headers))
    for j in range(18, len(headers) + 1):
        ws.cell(row=hr, column=j).fill = RUN_HDR_FILL

    for k, m in enumerate(meta, 1):
        r = hr + k
        s = f"'{m['sheet']}'"
        last = m["n"] + 1
        col = lambda name: f"{s}!${m['idx'][name]}$2:${m['idx'][name]}${last}"
        exp, act, lang = col("expected_result"), col("actual_result"), col("language")
        values = [
            k, m["validator"], m["sheet"], m["group"], m["origin"], m["meaning"], m["stage"], m["metric"], m["runner"],
            f"=COUNTA({col('case_id')})",
            f'=COUNTIF({exp},"FAIL")',
            f'=COUNTIF({exp},"PASS")',
            f"=IF(J{r}=0,0,K{r}/J{r})",
            f'=COUNTIF({lang},"en")',
            f'=COUNTIF({lang},"hi")',
            f'=COUNTIF({lang},"hi-Latn")',
            m["configs"],
            f'=COUNTIF({act},"PASS")+COUNTIF({act},"FAIL")',
            f'=COUNTIFS({exp},"FAIL",{act},"FAIL")',
            f'=COUNTIFS({exp},"PASS",{act},"FAIL")',
            f'=COUNTIFS({exp},"FAIL",{act},"PASS")',
            f'=COUNTIFS({exp},"PASS",{act},"PASS")',
            f'=COUNTIF({act},"ERROR")',
            f'=IF(R{r}=0,"",(S{r}+V{r})/R{r})',
            f'=IF(S{r}+T{r}=0,"",S{r}/(S{r}+T{r}))',
            f'=IF(S{r}+U{r}=0,"",S{r}/(S{r}+U{r}))',
            f'=IF(OR(Y{r}="",Z{r}=""),"",IF(Y{r}+Z{r}=0,0,2*Y{r}*Z{r}/(Y{r}+Z{r})))',
        ]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BASE
            cell.alignment = WRAP if c in (6, 8, 9) else TOP
            if c == 3:
                cell.hyperlink = f"#'{m['sheet']}'!A1"
                cell.font = Font(name=FONT, size=10, color="0563C1", underline="single")
            if c in (13, 24, 25, 26, 27):
                cell.number_format = "0.0%" if c == 13 else "0.00"

    tr = hr + len(meta) + 1
    ws.cell(row=tr, column=2, value="TOTAL").font = BOLD
    for c in list(range(10, 13)) + list(range(14, 17)) + list(range(18, 24)):
        L = get_column_letter(c)
        cell = ws.cell(row=tr, column=c, value=f"=SUM({L}{hr + 1}:{L}{tr - 1})")
        cell.font = BOLD
    ws.cell(row=tr, column=13, value=f"=IF(J{tr}=0,0,K{tr}/J{tr})").number_format = "0.0%"

    widths = [4, 26, 22, 15, 11, 34, 13, 30, 34, 7, 7, 7, 8, 7, 7, 8, 26, 9, 6, 6, 6, 6, 7, 9, 9, 9, 9]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=hr + 1, column=3)
    ws.row_dimensions[hr].height = 32


def write_configs(wb, rows):
    ws = wb.create_sheet("Configs", 2)
    ws["A1"] = "Validator configurations referenced by config_ref"
    ws["A1"].font = TITLE
    headers = ["config_ref", "validator", "parameters (JSON)", "description", "source file"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=3, column=c, value=h)
    style_header(ws, 3, len(headers))
    for i, (ref, v, params, desc, src) in enumerate(rows, 4):
        for c, val in enumerate([ref, v, json.dumps(params, indent=2, ensure_ascii=False), desc, src], 1):
            cell = ws.cell(row=i, column=c, value=val)
            cell.font = Font(name="Courier New", size=9) if c == 3 else BASE
            cell.alignment = WRAP
    for c, w in zip("ABCDE", [16, 26, 90, 45, 40]):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A4"


def write_run_log(wb, meta, history):
    ws = wb.create_sheet("Run_Log")
    ws["A1"] = "Run log — one row per validator per periodic run"
    ws["A1"].font = TITLE
    ws["A2"] = "Yellow cells are inputs. Copy the metric values from Index after filling actual_result. Rebuilding the workbook keeps these rows."
    ws["A2"].font = Font(name=FONT, italic=True, size=9)
    headers = ["run_date", "git_sha", "validator", "config_ref", "model / version", "rows run",
               "accuracy", "precision", "recall", "F1", "latency p50 (ms)", "latency p95 (ms)", "run by", "notes"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=4, column=c, value=h)
    style_header(ws, 4, len(headers))
    if history:
        for r, row in enumerate(history, 5):
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=v).font = BASE
    else:
        example = [date.today().isoformat(), "20e3f4a", "topic_relevance_llm", "TR-EDU", "gpt-5.6-sol", 192,
                   0.89, 0.86, 0.78, 0.82, 2700, 4100, "", "EXAMPLE ROW — replace"]
        for c, v in enumerate(example, 1):
            cell = ws.cell(row=5, column=c, value=v)
            cell.font = Font(name=FONT, size=10, italic=True, color="7F7F7F")
    for r in range(5, max(60, 5 + len(history) + 50)):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).fill = RUN_FILL
            if 7 <= c <= 10:
                ws.cell(row=r, column=c).number_format = "0.00"
    dv = DataValidation(type="list", formula1='"' + ",".join(m["validator"] for m in meta) + '"', allow_blank=True)
    dv.add(f"C5:C{max(59, 4 + len(history) + 50)}")
    ws.add_data_validation(dv)
    for c, w in enumerate([12, 10, 26, 14, 16, 9, 9, 9, 9, 9, 12, 12, 12, 40], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A5"


def read_run_log_history():
    """Existing Run_Log rows, so rebuilding the workbook keeps run history."""
    if not OUT.exists():
        return []
    ws = load_workbook(OUT)["Run_Log"]
    rows = [list(r) for r in ws.iter_rows(min_row=5, max_col=14, values_only=True) if r[0]]
    return [r for r in rows if not str(r[13] or "").startswith("EXAMPLE")]


def write_readme(wb, meta):
    ws = wb.active
    ws.title = "README"
    lines = [
        ("Guardrails validator evaluation set", TITLE),
        (f"Generated {date.today().isoformat()} by Benchmarking datasets/build_eval_workbook.py — do not hand-edit data sheets; edit the source CSVs and rebuild.", Font(name=FONT, italic=True, size=9)),
        ("", BASE),
        ("Purpose", BOLD),
        ("One standalone, periodically re-runnable evaluation set covering all 12 validators in validators.json. "
         "Every validator sheet shares one schema so the same runner/metrics logic applies everywhere.", BASE),
        ("", BASE),
        ("Sheets", BOLD),
        ("Index — one row per validator: data origin, coverage counts, and live confusion-matrix metrics (formulas).", BASE),
        ("Configs — the exact validator parameters each config_ref stands for (topic configs, ban lists, critic rubric).", BASE),
        ("<validator> sheets (12) — the test cases.", BASE),
        ("Run_Log — append one row per validator per periodic run to track drift over time.", BASE),
        ("", BASE),
        ("Column dictionary (all validator sheets)", BOLD),
        ("case_id — stable id, <PREFIX>-NNNN. Do not renumber between runs; append new cases at the end.", BASE),
        ("validator — validator type as in validators.json.", BASE),
        ("input_text — exact string passed to the validator. For answer_relevance it is the JSON {\"query\", \"answer\"} the validator expects.", BASE),
        ("expected_result — ground truth. FAIL = the validator SHOULD trigger (unsafe / PII / biased / out of scope / irrelevant). PASS = it should not.", BASE),
        ("label — 1 if FAIL else 0 (for scripts computing binary metrics).", BASE),
        ("expected_output — expected safe_text with on_fail=fix, where the validator has a fix (pii_remover, gender_assumption_bias, uli_slur_match, ban_list). Blank otherwise.", BASE),
        ("category — fine-grained case type; includes hard negatives (benign text that looks risky). Filter on it for error analysis.", BASE),
        ("difficulty — easy / medium / hard (synthetic sets only; blank for benchmark rows).", BASE),
        ("language — en | hi (Devanagari) | hi-Latn (romanized Hindi / Hinglish). For topic_relevance_llm rows it is auto-detected by a keyword heuristic.", BASE),
        ("domain — healthcare / education / agriculture / general / …", BASE),
        ("config_ref — key into the Configs sheet.", BASE),
        ("source — benchmark:<dataset> or synthetic.", BASE),
        ("notes — why a hard case is labelled the way it is.", BASE),
        ("Extra columns after notes are validator-specific (e.g. pii_entities, policy, banned_term_hit, failing_metrics, query/answer).", BASE),
        ("actual_result / actual_output (yellow) — fill from a run: PASS, FAIL or ERROR. match is a formula (Y/N).", BASE),
        ("", BASE),
        ("How to run periodically", BOLD),
        ("1. For each validator sheet, run the validator with the config named in config_ref over input_text (runners listed on Index).", BASE),
        ("2. Paste the outcome into actual_result (FAIL if the validator triggered) and the returned safe_text into actual_output.", BASE),
        ("3. Read accuracy / precision / recall / F1 off the Index sheet, then append one Run_Log row per validator.", BASE),
        ("4. Review rows where match = N, grouped by category — that is the error analysis for the report.", BASE),
        ("", BASE),
        ("Known data caveats", BOLD),
        ("gender_assumption_bias: benchmark has only 15 unique biased sentences repeated 20x. PASS rows are their neutral rewrites (as in the eval runner), so false positives are only measured on those 15 sentences.", BASE),
        ("content_safety: ~99% of the 601 'lexical' rows contain a term from the curated slur list, so report metrics per source_dataset — pooled recall is inflated for list-based checks.", BASE),
        ("content_safety: language comes from the source datasets; Hindi rows without Devanagari script are tagged hi-Latn. Some 'English' source rows are actually Hinglish.", BASE),
        ("ban_list: the hub BanList fix was off by one ('a sonography' -> 'a s'); fixed by app/core/validators/ban_list.py. expected_output is the ideal removal; compare after collapsing whitespace.", BASE),
        ("uli_slur_match: the slur list was not NFKC-normalised, so 12 Devanagari entries with nukta letters (e.g. बत्तमीज़, भड़वा) never matched; fixed in lexical_slur.py. Rows the list still misses (obfuscations, words not in the list) are FAIL with empty expected_output.", BASE),
        ("Full findings: Benchmarking datasets/FINDINGS.md.", BASE),
        ("Synthetic sets are LLM-written and spot-reviewed; treat hard-case labels as a starting point and refine with partner NGO data.", BASE),
    ]
    for i, (text, font) in enumerate(lines, 1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 140


def main():
    meta, frames = [], []
    for validator, sheet, prefix, group, origin, loader, meaning, metric, stage, runner in VALIDATORS:
        df = loader() if loader else load_synthetic(validator, prefix)
        frames.append((sheet, df))
        meta.append(dict(validator=validator, sheet=sheet, group=group, origin=origin, meaning=meaning,
                         metric=metric, stage=stage, runner=runner, n=len(df),
                         configs=", ".join(sorted(df["config_ref"].unique()))))

    wb = Workbook()
    write_readme(wb, meta)
    for (sheet, df), m in zip(frames, meta):
        idx, _ = write_validator_sheet(wb, sheet, df)
        m["idx"] = {name: get_column_letter(i) for name, i in idx.items()}
    write_index(wb, meta)
    write_configs(wb, build_configs())
    write_run_log(wb, meta, read_run_log_history())
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    for m in meta:
        print(f"{m['validator']:<28} {m['origin']:<10} {m['n']:>5} rows")
    print(f"saved {OUT}")


if __name__ == "__main__":
    main()
